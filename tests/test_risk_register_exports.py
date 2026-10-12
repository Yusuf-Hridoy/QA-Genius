"""The risk register as CSV and Excel. No network."""

import csv
import io

import openpyxl

from qagenius.risk_radar import score_risks
from qagenius.risk_register_exports import COLUMNS, HEAT_SHEET, REGISTER_SHEET, to_csv, to_xlsx
from tests.factories import case
from tests.test_risk_radar import analysis, mode


def _report(*modes):
    return score_risks(analysis(*modes), [case(case_id="TC-001")])


def _rows(data: bytes) -> list[list[str]]:
    text = data.decode("utf-8")
    assert text.startswith("\ufeff")
    return list(csv.reader(io.StringIO(text[1:])))


def test_csv_starts_with_a_bom_and_the_header() -> None:
    rows = _rows(to_csv(_report(mode("R1"), mode("R2", title="Lock email delayed"))))
    assert rows[0] == list(COLUMNS)
    assert len(rows) == 3
    assert rows[1][0] == "R1"


def test_csv_keeps_the_computed_numbers() -> None:
    rows = _rows(to_csv(_report(mode("R1", severity=5, likelihood=4, detectability=4))))
    header = list(COLUMNS)
    row = rows[1]
    assert row[header.index("RPN")] == "80"
    assert row[header.index("Level")] == "high"
    assert row[header.index("Coverage")] == "none"


def test_a_title_that_looks_like_a_formula_is_defused() -> None:
    rows = _rows(to_csv(_report(mode("R1", title="=SUM(A1)"))))
    assert rows[1][1] == "'=SUM(A1)"


def test_the_workbook_has_both_sheets() -> None:
    book = openpyxl.load_workbook(io.BytesIO(to_xlsx(_report(mode("R1")))))
    assert book.sheetnames == [REGISTER_SHEET, HEAT_SHEET]
    sheet = book[REGISTER_SHEET]
    assert [cell.value for cell in sheet[1]] == list(COLUMNS)
    assert sheet.freeze_panes == "A2"
    assert sheet[1][0].font.bold


def test_the_level_cell_is_filled_by_level() -> None:
    report = _report(
        mode("R1", severity=5, likelihood=4, detectability=4),  # high
        mode("R2", title="Medium one", severity=4, likelihood=3, detectability=3),
        mode("R3", title="Low one", severity=2, likelihood=2, detectability=2),
    )
    sheet = openpyxl.load_workbook(io.BytesIO(to_xlsx(report)))[REGISTER_SHEET]
    column = list(COLUMNS).index("Level") + 1
    fills = {
        sheet.cell(row=row, column=1).value: sheet.cell(row=row, column=column).fill.fgColor.rgb
        for row in range(2, 5)
    }
    assert fills["R1"].endswith("F7DEDA")
    assert fills["R2"].endswith("F6E9CF")
    assert fills["R3"].endswith("DDEBE3")


def test_the_heat_map_sheet_puts_the_risk_in_its_cell() -> None:
    report = _report(mode("R1", severity=5, likelihood=4))
    sheet = openpyxl.load_workbook(io.BytesIO(to_xlsx(report)))[HEAT_SHEET]
    assert [cell.value for cell in sheet[1]][1:] == ["L1", "L2", "L3", "L4", "L5"]
    # Row 2 is severity 5; column E is likelihood 4.
    assert sheet.cell(row=2, column=1).value == "S5"
    assert sheet.cell(row=2, column=5).value == "R1"
    assert sheet.cell(row=6, column=1).value == "S1"
    # An empty cell: openpyxl reads a blank string back as None.
    assert not sheet.cell(row=3, column=2).value
