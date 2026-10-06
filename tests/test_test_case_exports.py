"""Export tests. No network: the suite is built by hand."""

import csv
import io

import openpyxl

# Aliased: pytest tries to collect any module-level name starting with 'Test'.
from qagenius.models import TestCase as Case
from qagenius.models import TestCaseList as CaseList
from qagenius.models import TestSuiteSummary as Summary
from qagenius.test_case_exports import COLUMNS, _safe_cell, to_csv, to_xlsx
from qagenius.test_cases import number_criteria

CRITERIA = number_criteria(
    [
        "The account locks after 5 failed attempts",
        "The lock lasts 30 minutes",
        "A locked shopper sees a message",
    ]
)


def _case(
    case_id: str = "TC-001",
    traceability: str = "AC-1",
    title: str = "Lock after five wrong passwords",
    test_data: str | None = "email=shopper@example.com",
    bdd: str | None = "Given a shopper\nWhen they fail 5 times\nThen the account locks",
) -> Case:
    return Case(
        id=case_id,
        title=title,
        category="Functional",
        pre_conditions="The shopper has an active Aurora Storefront account",
        steps=["Open the login page", "Enter a wrong password five times"],
        expected_result="The account locks for 30 minutes",
        priority="High",
        test_data=test_data,
        bdd_scenario=bdd,
        automation_feasibility="High",
        automation_effort="Low",
        tags=["login", "security"],
        traceability=traceability,
    )


def _suite(*cases: Case) -> CaseList:
    return CaseList(
        test_cases=list(cases),
        summary=Summary(
            total_generated=999,
            category_breakdown=[],
            priority_breakdown=[],
            automation_coverage_potential="99%",
            coverage_gaps=[],
            recommendations=[],
        ),
    )


def test_csv_starts_with_bom_and_header() -> None:
    data = to_csv(_suite(_case()), CRITERIA)
    assert data.startswith(b"\xef\xbb\xbf")
    text = data.decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(text)))
    assert rows[0] == list(COLUMNS)


def test_csv_row_count_and_traceability() -> None:
    suite = _suite(
        _case("TC-001", traceability="AC-1"),
        _case("TC-002", traceability="Covers AC-2 and AC-3"),
        _case("TC-003", traceability="AC-9"),
    )
    rows = list(csv.reader(io.StringIO(to_csv(suite, CRITERIA).decode("utf-8-sig"))))
    assert len(rows) == 4  # 3 cases + header
    covers = COLUMNS.index("Covers (AC)")
    assert rows[1][covers] == "AC-1"
    assert rows[2][covers] == "AC-2, AC-3"
    assert rows[3][covers] == ""  # AC-9 is not a real criterion


def test_csv_steps_are_numbered() -> None:
    rows = list(csv.reader(io.StringIO(to_csv(_suite(_case()), CRITERIA).decode("utf-8-sig"))))
    steps = rows[1][COLUMNS.index("Steps")]
    assert steps == "1. Open the login page\n2. Enter a wrong password five times"


def test_csv_missing_optional_fields_are_blank() -> None:
    suite = _suite(_case(test_data=None, bdd=None))
    rows = list(csv.reader(io.StringIO(to_csv(suite, CRITERIA).decode("utf-8-sig"))))
    assert rows[1][COLUMNS.index("Test Data")] == ""
    assert rows[1][COLUMNS.index("BDD Scenario")] == ""


def test_safe_cell_blocks_formulas() -> None:
    assert _safe_cell("=SUM(A1)") == "'=SUM(A1)"
    assert _safe_cell("-1 day") == "'-1 day"
    assert _safe_cell("+1") == "'+1"
    assert _safe_cell("@here") == "'@here"
    assert _safe_cell("Lock the account") == "Lock the account"


def test_csv_escapes_formula_in_title() -> None:
    rows = list(
        csv.reader(
            io.StringIO(
                to_csv(_suite(_case(title="=cmd|' /c calc'!A1")), CRITERIA).decode("utf-8-sig")
            )
        )
    )
    assert rows[1][COLUMNS.index("Title")].startswith("'=")


def test_xlsx_sheets_and_header() -> None:
    book = openpyxl.load_workbook(io.BytesIO(to_xlsx(_suite(_case()), CRITERIA)))
    assert book.sheetnames == ["Test Cases", "Coverage"]
    sheet = book["Test Cases"]
    assert [cell.value for cell in sheet[1]] == list(COLUMNS)
    assert all(cell.font.bold for cell in sheet[1])
    assert sheet.freeze_panes == "A2"


def test_xlsx_coverage_sheet() -> None:
    suite = _suite(_case("TC-001", traceability="AC-1"), _case("TC-002", traceability="AC-1"))
    book = openpyxl.load_workbook(io.BytesIO(to_xlsx(suite, CRITERIA)))
    sheet = book["Coverage"]
    assert [cell.value for cell in sheet[1]] == ["Criterion", "Text", "Test cases"]
    assert sheet["A2"].value == "AC-1"
    assert sheet["C2"].value == 2
    # AC-2 and AC-3 are uncovered, so their count cells are flagged.
    assert sheet["C3"].value == 0
    assert sheet["C3"].fill.fgColor.rgb.endswith("F7DEDA")
    assert sheet["C2"].fill.fgColor.rgb != sheet["C3"].fill.fgColor.rgb
    last = sheet.max_row
    assert sheet.cell(row=last, column=1).value == "Coverage"
    assert sheet.cell(row=last, column=2).value == "33%"


def test_xlsx_escapes_formula() -> None:
    suite = _suite(_case(title="=HYPERLINK('http://x')"))
    book = openpyxl.load_workbook(io.BytesIO(to_xlsx(suite, CRITERIA)))
    assert book["Test Cases"]["B2"].value.startswith("'=")
