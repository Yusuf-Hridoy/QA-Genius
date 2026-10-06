"""Write a generated test suite out as CSV or Excel, with a coverage sheet."""

import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from qagenius.models import TestCase, TestCaseList
from qagenius.test_cases import Criterion, compute_coverage, traced_ids

COLUMNS = (
    "ID",
    "Title",
    "Category",
    "Priority",
    "Covers (AC)",
    "Pre-conditions",
    "Steps",
    "Expected Result",
    "Test Data",
    "BDD Scenario",
    "Automation Feasibility",
    "Automation Effort",
    "Tags",
)

COVERAGE_COLUMNS = ("Criterion", "Text", "Test cases")

# Excel and Calc read these as the start of a formula.
_FORMULA_STARTS = ("=", "+", "-", "@")

_WIDTHS = {"ID": 10, "Title": 40, "Steps": 60, "Expected Result": 50}
_DEFAULT_WIDTH = 18
_WRAP_COLUMNS = ("Steps", "Expected Result", "BDD Scenario")
_UNCOVERED_FILL = PatternFill("solid", fgColor="F7DEDA")


def _safe_cell(value: str) -> str:
    """Stop a spreadsheet reading a cell as a formula."""
    text = "" if value is None else str(value)
    if text.startswith(_FORMULA_STARTS):
        return "'" + text
    return text


def _numbered_steps(steps: list[str]) -> str:
    return "\n".join(f"{i}. {step}" for i, step in enumerate(steps, start=1))


def _row(test_case: TestCase, valid_ids: set[str]) -> list[str]:
    values = [
        test_case.id,
        test_case.title,
        test_case.category,
        test_case.priority,
        ", ".join(traced_ids(test_case, valid_ids)),
        test_case.pre_conditions,
        _numbered_steps(test_case.steps),
        test_case.expected_result,
        test_case.test_data or "",
        test_case.bdd_scenario or "",
        test_case.automation_feasibility,
        test_case.automation_effort,
        "; ".join(test_case.tags or []),
    ]
    return [_safe_cell(value) for value in values]


def to_csv(result: TestCaseList, criteria: list[Criterion]) -> bytes:
    """UTF-8 with a BOM, so Excel opens it without an import step."""
    valid_ids = {criterion.id for criterion in criteria}
    buffer = io.StringIO()
    writer = csv.writer(buffer, quoting=csv.QUOTE_MINIMAL, lineterminator="\r\n")
    writer.writerow(COLUMNS)
    for test_case in result.test_cases:
        writer.writerow(_row(test_case, valid_ids))
    return ("﻿" + buffer.getvalue()).encode("utf-8")


def _write_cases_sheet(sheet, result: TestCaseList, valid_ids: set[str]) -> None:
    sheet.append(list(COLUMNS))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    for index, name in enumerate(COLUMNS, start=1):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = _WIDTHS.get(name, _DEFAULT_WIDTH)
    for test_case in result.test_cases:
        sheet.append(_row(test_case, valid_ids))
    wrap = Alignment(wrap_text=True, vertical="top")
    for name in _WRAP_COLUMNS:
        letter = get_column_letter(COLUMNS.index(name) + 1)
        for row in range(2, sheet.max_row + 1):
            sheet[f"{letter}{row}"].alignment = wrap


def _write_coverage_sheet(sheet, result: TestCaseList, criteria: list[Criterion]) -> None:
    coverage = compute_coverage(result, criteria)
    sheet.append(list(COVERAGE_COLUMNS))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.column_dimensions["A"].width = 14
    sheet.column_dimensions["B"].width = 70
    sheet.column_dimensions["C"].width = 12
    for criterion in criteria:
        count = coverage.cases_per_id[criterion.id]
        sheet.append([_safe_cell(criterion.id), _safe_cell(criterion.text), count])
        if count == 0:
            sheet.cell(row=sheet.max_row, column=3).fill = _UNCOVERED_FILL
    sheet.append(["Coverage", f"{coverage.percent}%"])


def to_xlsx(result: TestCaseList, criteria: list[Criterion]) -> bytes:
    """A workbook with the cases and a coverage sheet that flags the gaps."""
    valid_ids = {criterion.id for criterion in criteria}
    book = Workbook()
    cases_sheet = book.active
    cases_sheet.title = "Test Cases"
    _write_cases_sheet(cases_sheet, result, valid_ids)
    _write_coverage_sheet(book.create_sheet("Coverage"), result, criteria)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
